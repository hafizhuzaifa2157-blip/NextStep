
from flask import Flask, request, render_template
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document
from io import BytesIO
from collections import Counter
import re
import requests

app = Flask(__name__)


# -------------------- HOME AND MAIN PAGES --------------------

@app.route("/")
def home():
    return render_template("index.html")


@app.route("/profile")
def profile():
    return render_template("profile.html")


@app.route("/dashboard")
def dashboard():
    return render_template("dashboard.html")


# -------------------- CV BUILDER --------------------

@app.route("/cv-builder")
def cv_builder():
    return render_template("cv_builder.html")


# -------------------- CV TEXT EXTRACTION --------------------

def extract_cv_text(cv_file):
    """Extract readable text from PDF, DOCX or TXT."""

    filename = (cv_file.filename or "").lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(cv_file.stream)

        if reader.is_encrypted:
            raise ValueError("Please upload an unlocked PDF.")

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif filename.endswith(".docx"):
        document = Document(cv_file.stream)

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " ".join(cell.text for cell in row.cells)
                )

        return "\n".join(parts)

    elif filename.endswith(".txt"):
        return cv_file.read().decode(
            "utf-8-sig", errors="ignore"
        )

    raise ValueError(
        "Please upload a PDF, DOCX, or TXT file."
    )


# -------------------- CV ANALYZER --------------------

@app.route("/cv-analyzer", methods=["GET", "POST"])
def cv_analyzer():
    return render_template("cv_analyzer.html")


# -------------------- ATS-FRIENDLY CV CHECK --------------------

@app.route("/ats-check", methods=["GET", "POST"])
def ats_check():
    result = None
    error = None

    if request.method == "POST":
        cv_file = request.files.get("cv_file")
        job_description = request.form.get(
            "job_description", ""
        ).strip()

        if not cv_file or not cv_file.filename:
            error = "Please select your CV file."

        elif not job_description:
            error = "Please enter the job description."

        else:
            try:
                cv_text = extract_cv_text(cv_file)

                if not cv_text.strip():
                    error = (
                        "No readable text found. Please try "
                        "a text-based PDF, DOCX, or TXT file."
                    )
                else:
                    cv_lower = cv_text.lower()
                    job_lower = job_description.lower()

                    keywords = [
                        "python", "flask", "django", "sql",
                        "mysql", "postgresql", "rest api",
                        "api", "git", "github", "html", "css",
                        "javascript", "java", "communication",
                        "teamwork", "problem solving",
                        "data analysis", "machine learning"
                    ]

                    job_keywords = list(dict.fromkeys(
                        word for word in keywords
                        if word in job_lower
                    ))

                    matched = [
                        word for word in job_keywords
                        if word in cv_lower
                    ]

                    missing = [
                        word for word in job_keywords
                        if word not in cv_lower
                    ]

                    keyword_score = (
                        round(
                            len(matched) / len(job_keywords) * 60
                        )
                        if job_keywords else 0
                    )

                    sections = {
                        "Education": any(word in cv_lower for word in [
                            "education", "academic", "school",
                            "grade 9", "class 9", "matric"
                        ]),
                        "Experience": any(word in cv_lower for word in [
                            "experience", "work history",
                            "volunteer", "volunteering", "activities"
                        ]),
                        "Skills": any(word in cv_lower for word in [
                            "skills", "technical skills",
                            "python", "html", "css"
                        ]),
                        "Projects": any(word in cv_lower for word in [
                            "projects", "project"
                        ])
                    }

                    section_score = sum(sections.values()) * 10

                    score = min(
                        keyword_score + section_score, 100
                    )

                    result = {
                        "score": score,
                        "matched_keywords": matched,
                        "missing_keywords": missing,
                        "sections": sections
                    }

            except ValueError as exc:
                error = str(exc)

            except Exception:
                error = (
                    "Could not read this file. Please try "
                    "another PDF, DOCX, or TXT file."
                )

    return render_template(
        "ats_check.html",
        result=result,
        error=error
    )


# -------------------- CV IMPROVEMENT --------------------

@app.route("/cv-improvement", methods=["GET", "POST"])
def cv_improvement():
    suggestions = []
    error = None

    if request.method == "POST":
        cv_text = request.form.get("cv_text", "").strip()
        cv_file = request.files.get("cv_file")

        if cv_file and cv_file.filename:
            try:
                cv_text = extract_cv_text(cv_file)

            except ValueError as exc:
                error = str(exc)

            except Exception:
                error = (
                    "Could not read this file. Please try "
                    "another PDF, DOCX, or TXT file."
                )

        if not error:
            if not cv_text.strip():
                error = (
                    "Please upload your CV or paste its text."
                )
            else:
                text = cv_text.lower()

                checks = {
                    "email": "@" in text,
                    "education": any(word in text for word in [
                        "education", "academic", "school",
                        "grade 9", "class 9", "matric"
                    ]),
                    "skills": any(word in text for word in [
                        "skills", "technical skills",
                        "python", "html", "css"
                    ]),
                    "projects": any(word in text for word in [
                        "projects", "project",
                        "student management", "parking management"
                    ]),
                    "experience": any(word in text for word in [
                        "experience", "work history",
                        "volunteer", "volunteering", "activities"
                    ])
                }

                if not checks["email"]:
                    suggestions.append(
                        "Add a professional email address."
                    )

                if not checks["education"]:
                    suggestions.append(
                        "Add your education level and school details."
                    )

                if not checks["skills"]:
                    suggestions.append(
                        "Add a clear Skills section with relevant skills."
                    )

                if not checks["projects"]:
                    suggestions.append(
                        "Add a Projects section describing your work."
                    )

                if not checks["experience"]:
                    suggestions.append(
                        "Add relevant experience, school activities, "
                        "or volunteer work."
                    )

                if len(cv_text.split()) < 100:
                    suggestions.append(
                        "Add more relevant details about your "
                        "education, skills, projects, and achievements."
                    )

                if not suggestions:
                    suggestions = [
                        "Your CV passed these basic section checks.",
                        "Explain what you built in each project "
                        "and which tools you used.",
                        "Check spelling, grammar, and readability.",
                        "Tailor your skills to the job description.",
                        "Include genuine achievements where relevant."
                    ]

    return render_template(
        "cv_improvement.html",
        suggestions=suggestions,
        error=error
    )


# -------------------- JOB RECOMMENDATIONS --------------------

@app.route("/job-recommendations", methods=["GET", "POST"])
def job_recommendations():
    jobs = []
    career = ""
    level = ""
    error = ""

    if request.method == "POST":
        career = request.form.get("career", "").strip()
        level = request.form.get("level", "").strip()

        if career and level:
            try:
                response = requests.get(
                    "https://remotive.com/api/remote-jobs",
                    params={"search": career},
                    timeout=10
                )
                response.raise_for_status()
                data = response.json()

                for job in data.get("jobs", [])[:15]:
                    jobs.append({
                        "title": job.get(
                            "title", "Title not provided"
                        ),
                        "company": job.get(
                            "company_name", "Company not provided"
                        ),
                        "location": job.get(
                            "candidate_required_location", "Remote"
                        ),
                        "salary": job.get("salary") or "Not specified",
                        "url": job.get("url", "#"),
                        "description": job.get("description", ""),
                        "level": level
                    })

            except Exception:
                error = (
                    "Jobs could not be loaded. Please try again."
                )

    return render_template(
        "job_recommendations.html",
        jobs=jobs,
        career=career,
        level=level,
        error=error
    )


# -------------------- JOB MATCHING --------------------

@app.route("/job-matching", methods=["GET", "POST"])
def job_matching():
    score = None
    matched = []
    missing = []
    error = None
    job_description = ""

    if request.method == "POST":
        uploaded_file = request.files.get("cv_file")
        job_description = request.form.get(
            "job_description", ""
        ).strip()

        if not uploaded_file or not uploaded_file.filename:
            error = "Please upload your CV."

        elif not job_description:
            error = "Please enter the job description."

        else:
            try:
                filename = secure_filename(
                    uploaded_file.filename
                )
                extension = filename.rsplit(".", 1)[-1].lower()
                file_data = uploaded_file.read()

                if extension == "pdf":
                    reader = PdfReader(BytesIO(file_data))

                    if reader.is_encrypted:
                        raise ValueError(
                            "Please upload an unlocked PDF."
                        )

                    cv_text = "\n".join(
                        page.extract_text() or ""
                        for page in reader.pages
                    )

                elif extension == "docx":
                    document = Document(BytesIO(file_data))

                    parts = [
                        paragraph.text
                        for paragraph in document.paragraphs
                    ]

                    for table in document.tables:
                        for row in table.rows:
                            parts.append(
                                " ".join(
                                    cell.text for cell in row.cells
                                )
                            )

                    cv_text = "\n".join(parts)

                elif extension == "txt":
                    cv_text = file_data.decode(
                        "utf-8-sig", errors="ignore"
                    )

                else:
                    raise ValueError(
                        "Only PDF, DOCX and TXT files are supported."
                    )

                if not cv_text.strip():
                    raise ValueError(
                        "No readable text was found in your CV. "
                        "Please upload a text-based PDF, DOCX or TXT file."
                    )

                stop_words = {
                    "the", "and", "for", "with", "that", "this",
                    "you", "your", "are", "our", "from", "will",
                    "have", "has", "who", "job", "work", "working",
                    "about", "into", "their", "they", "can", "all",
                    "not", "but", "per", "any", "one"
                }

                keywords = sorted({
                    word.strip(".,:;!?()[]{}\"'")
                    for word in job_description.lower().split()
                    if len(word.strip(".,:;!?()[]{}\"'")) > 2
                    and word.strip(".,:;!?()[]{}\"'") not in stop_words
                })

                cv_lower = cv_text.lower()

                matched = [
                    keyword for keyword in keywords
                    if keyword in cv_lower
                ]

                missing = [
                    keyword for keyword in keywords
                    if keyword not in cv_lower
                ]

                if keywords:
                    score = round(
                        len(matched) / len(keywords) * 100
                    )
                else:
                    error = (
                        "The job description does not contain "
                        "enough useful keywords to calculate a score."
                    )

            except Exception as exc:
                error = str(exc) or (
                    "Unable to read this file. Please try another CV."
                )

    return render_template(
        "job_matching.html",
        score=score,
        matched=matched,
        missing=missing,
        error=error,
        job_description=job_description
    )


# -------------------- JOB DESCRIPTION ANALYZER --------------------

@app.route("/job-description-analyzer", methods=["GET", "POST"])
def job_description_analyzer():
    description = ""
    keywords = []
    experience = "Not specified"
    education = "Not specified"

    if request.method == "POST":
        description = request.form.get(
            "description", ""
        ).strip()

        if description:
            words = re.findall(
                r"[a-zA-Z][a-zA-Z+#.]*",
                description.lower()
            )

            stop_words = {
                "the", "and", "for", "with", "you", "your",
                "our", "are", "will", "this", "that", "from",
                "have", "has", "who", "job", "work", "about",
                "years", "year", "ability", "skills", "experience",
                "knowledge", "strong", "good", "required", "preferred"
            }

            keywords = [
                word for word, count in Counter(words).most_common(30)
                if word not in stop_words and len(word) > 2
            ][:15]

            match = re.search(
                r"(\d+\+?\s*(?:to|-)\s*\d+\s*years?|\d+\+?\s*years?)",
                description.lower()
            )

            if match:
                experience = match.group(0)

            for line in description.splitlines():
                if any(term in line.lower() for term in [
                    "bachelor", "master", "degree", "diploma",
                    "qualification", "graduate", "education"
                ]):
                    education = line.strip()[:200]
                    break

    return render_template(
        "job_description_analyzer.html",
        description=description,
        keywords=keywords,
        experience=experience,
        education=education
    )


# -------------------- MISSING KEYWORDS ANALYZER --------------------

@app.route("/missing-keywords-analyzer", methods=["GET", "POST"])
def missing_keywords_analyzer():
    missing = []
    matched = []
    error = None
    job_description = ""

    if request.method == "POST":
        uploaded_file = request.files.get("cv_file")
        job_description = request.form.get(
            "job_description", ""
        ).strip()

        try:
            if not uploaded_file or not uploaded_file.filename:
                raise ValueError("Please upload your CV.")

            if not job_description:
                raise ValueError(
                    "Please enter the job description."
                )

            cv_text = extract_cv_text(uploaded_file)

            if not cv_text.strip():
                raise ValueError(
                    "No readable text found in your CV."
                )

            stop_words = {
                "the", "and", "for", "with", "you", "your",
                "our", "are", "will", "this", "that", "from",
                "have", "has", "who", "job", "work", "about",
                "years", "year", "ability", "required", "preferred"
            }

            keywords = set(re.findall(
                r"[a-zA-Z][a-zA-Z+#.]*",
                job_description.lower()
            ))

            keywords = {
                word for word in keywords
                if len(word) > 2 and word not in stop_words
            }

            cv_lower = cv_text.lower()

            matched = sorted(
                word for word in keywords
                if word in cv_lower
            )

            missing = sorted(
                word for word in keywords
                if word not in cv_lower
            )

        except Exception as exc:
            error = str(exc) or "Unable to analyze this CV."

    return render_template(
        "missing_keywords_analyzer.html",
        matched=matched,
        missing=missing,
        error=error,
        job_description=job_description
    )


# -------------------- SKILL GAP ANALYSIS --------------------

@app.route("/skill-gap-analysis", methods=["GET", "POST"])
def skill_gap_analysis():
    current_skills = ""
    required_skills = ""
    present = []
    gaps = []
    roadmap = []
    progress = 0
    error = None

    skill_advice = {
        "python": (
            "High Priority",
            "Practise variables, conditions, loops, and functions. "
            "Then build a small Python project."
        ),
        "flask": (
            "High Priority",
            "Learn Flask routes, HTML templates, and form handling."
        ),
        "sql": (
            "High Priority",
            "Practise SQL SELECT, INSERT, UPDATE, and DELETE queries."
        ),
        "html": (
            "Medium Priority",
            "Practise HTML headings, links, forms, and page structure."
        ),
        "css": (
            "Medium Priority",
            "Learn CSS colours, spacing, Flexbox, and responsive design."
        ),
        "git": (
            "Medium Priority",
            "Learn repositories, commits, and version control."
        ),
        "javascript": (
            "Medium Priority",
            "Learn JavaScript variables, functions, and interactive features."
        ),
        "communication": (
            "Medium Priority",
            "Practise explaining and presenting your projects clearly."
        ),
        "problem solving": (
            "Medium Priority",
            "Practise breaking problems into smaller steps and solving them."
        )
    }

    aliases = {
        "python programming": "python",
        "python programming language": "python",
        "flask framework": "flask",
        "structured query language": "sql",
        "version control": "git",
        "problem-solving": "problem solving"
    }

    def normalize_skill(skill):
        value = " ".join(skill.lower().strip().split())
        return aliases.get(value, value)

    if request.method == "POST":
        current_skills = request.form.get(
            "current_skills", ""
        ).strip()

        required_skills = request.form.get(
            "required_skills", ""
        ).strip()

        if not current_skills or not required_skills:
            error = (
                "Please enter both your current skills "
                "and the skills required for your target job."
            )

        else:
            current = {
                normalize_skill(item)
                for item in re.split(r"[,;\n]+", current_skills)
                if item.strip()
            }

            required = []
            seen = set()

            for item in re.split(r"[,;\n]+", required_skills):
                item = item.strip()
                normalized = normalize_skill(item)

                if normalized and normalized not in seen:
                    required.append((item, normalized))
                    seen.add(normalized)

            for display_name, normalized in required:
                if normalized in current:
                    present.append(display_name)

                else:
                    priority, task = skill_advice.get(
                        normalized,
                        (
                            "Practice Needed",
                            f"Learn the basics of {display_name}, "
                            "practise regularly, and build a small project."
                        )
                    )

                    gaps.append({
                        "name": display_name,
                        "priority": priority,
                        "task": task
                    })

            if required:
                progress = round(
                    len(present) / len(required) * 100
                )

            roadmap = [
                {
                    "name": item["name"],
                    "task": item["task"]
                }
                for item in gaps
            ]

    return render_template(
        "skill_gap_analysis.html",
        current_skills=current_skills,
        required_skills=required_skills,
        present=present,
        gaps=gaps,
        roadmap=roadmap,
        progress=progress,
        error=error
    )


# -------------------- LEARNING ROADMAP --------------------

@app.route("/learning-roadmap", methods=["GET", "POST"])
def learning_roadmap():
    goal = ""
    job_description = ""
    current_skills = []
    missing_skills = []
    error = None
    cv_uploaded = False

    career_skills = {
        "web developer": {
            "html": "Practise webpage structure, links, and forms.",
            "css": "Learn styling, Flexbox, Grid, and responsive design.",
            "javascript": "Learn interactive webpages, functions, and events.",
            "python": "Practise Python fundamentals and build projects.",
            "flask": "Learn routes, templates, and form handling.",
            "git": "Learn version control and GitHub."
        },
        "software developer": {
            "python": "Practise functions, data structures, and programming.",
            "problem solving": "Break problems into smaller steps and solve them.",
            "git": "Learn version control and GitHub.",
            "sql": "Practise database queries and data management.",
            "testing": "Learn to test code and fix errors."
        },
        "python developer": {
            "python": "Improve Python fundamentals and data structures.",
            "flask": "Build web applications with Flask.",
            "sql": "Practise SQL and database operations.",
            "git": "Learn version control and GitHub.",
            "testing": "Write tests and practise debugging."
        },
        "data analyst": {
            "python": "Practise analysing and cleaning data.",
            "sql": "Learn queries, filtering, grouping, and joins.",
            "excel": "Practise formulas, sorting, and pivot tables.",
            "data visualization": "Present findings using clear charts.",
            "statistics": "Learn averages, percentages, and basic statistics."
        },
        "graphic designer": {
            "photoshop": "Practise image editing, layers, and visual composition.",
            "figma": "Learn interface design, components, and prototyping.",
            "communication": "Practise presenting and explaining design decisions."
        },
        "ui ux designer": {
            "figma": "Practise wireframes, prototypes, and interface design.",
            "communication": "Learn to explain design decisions and user needs.",
            "problem solving": "Practise identifying user problems and testing solutions."
        },
        "cybersecurity analyst": {
            "cybersecurity": "Study security fundamentals and safe lab exercises.",
            "networking": "Learn networking concepts and protocols.",
            "python": "Use Python for basic automation and data handling."
        }
    }

    skill_advice = {
        "python": "Practise Python and build a small project.",
        "html": "Build webpages using headings, links, and forms.",
        "css": "Practise styling and responsive layouts.",
        "javascript": "Practise JavaScript functions and events.",
        "flask": "Build a small Flask application.",
        "sql": "Practise SQL queries and database operations.",
        "git": "Practise saving projects with Git and GitHub.",
        "excel": "Practise spreadsheets, formulas, and data analysis.",
        "communication": "Practise explaining ideas and presenting your work clearly.",
        "problem solving": "Solve small problems step by step.",
        "testing": "Learn to test programs and fix errors.",
        "statistics": "Practise basic statistics and interpreting data.",
        "data visualization": "Create charts to communicate information.",
        "photoshop": "Practise image editing and visual design.",
        "figma": "Practise designing interfaces and prototypes.",
        "machine learning": "Learn data preparation and basic machine learning.",
        "cybersecurity": "Study security fundamentals and safe lab exercises.",
        "networking": "Learn networking concepts and protocols."
    }

    skill_aliases = {
        "js": "javascript",
        "html5": "html",
        "css3": "css",
        "github": "git",
        "mysql": "sql",
        "postgresql": "sql",
        "problem-solving": "problem solving",
        "ui/ux designer": "ui ux designer",
        "ui ux design": "ui ux designer",
        "front end developer": "web developer",
        "frontend developer": "web developer",
        "website developer": "web developer",
        "full stack developer": "web developer",
        "back end developer": "software developer",
        "backend developer": "software developer",
        "software engineer": "software developer",
        "python programmer": "python developer",
        "data scientist": "data analyst",
        "graphic design": "graphic designer",
        "cybersecurity": "cybersecurity analyst"
    }

    def normalize_career(value):
        value = " ".join(value.lower().strip().split())
        return skill_aliases.get(value, value)

    if request.method == "POST":
        goal = request.form.get("goal", "").strip()
        job_description = request.form.get(
            "job_description", ""
        ).strip()
        cv_file = request.files.get("cv_file")

        if not goal:
            error = "Please enter your target career."

        elif not cv_file or not cv_file.filename:
            error = "Please upload your CV."

        else:
            try:
                cv_text = extract_cv_text(cv_file)

                if not cv_text.strip():
                    error = (
                        "No readable text was found in your CV. "
                        "Please try another PDF, DOCX, or TXT file."
                    )

                else:
                    cv_uploaded = True
                    cv_lower = cv_text.lower()

                    # Identify skills mentioned in the CV.
                    for skill in skill_advice:
                        if re.search(
                            r"(?<!\w)" + re.escape(skill) + r"(?!\w)",
                            cv_lower
                        ):
                            current_skills.append(skill)

                    # Match the target career to our checklist.
                    matched_career = normalize_career(goal)
                    required = career_skills.get(matched_career)

                    # If the career is not in our list, inspect
                    # the job description for known skills.
                    if not required and job_description:
                        description_lower = job_description.lower()

                        required = {
                            skill: advice
                            for skill, advice in skill_advice.items()
                            if re.search(
                                r"(?<!\w)" + re.escape(skill) + r"(?!\w)",
                                description_lower
                            )
                            or any(
                                alias in description_lower
                                for alias, canonical in skill_aliases.items()
                                if canonical == skill
                            )
                        }

                    # Provide basic fallback suggestions when
                    # no known skills were found.
                    if not required:
                        required = {
                            "communication": (
                                "Practise explaining ideas clearly."
                            ),
                            "problem solving": (
                                "Practise solving problems step by step."
                            ),
                            "python": (
                                "Learn Python basics and build a small project."
                            ),
                            "git": (
                                "Learn Git and GitHub basics."
                            )
                        }

                    # Show the skills not mentioned in the CV.
                    for skill, advice in required.items():
                        if skill not in current_skills:
                            missing_skills.append({
                                "name": skill.title(),
                                "task": advice
                            })

            except ValueError as exc:
                error = str(exc)
                cv_uploaded = False

            except Exception:
                error = (
                    "Could not read your CV. Please upload a readable "
                    "PDF, DOCX, or TXT file."
                )
                cv_uploaded = False

    return render_template(
        "learning_roadmap.html",
        goal=goal,
        job_description=job_description,
        current_skills=current_skills,
        missing_skills=missing_skills,
        error=error,
        cv_uploaded=cv_uploaded
    )


# -------------------- RUN APPLICATION --------------------

if __name__ == "__main__":
    app.run(debug=True)
